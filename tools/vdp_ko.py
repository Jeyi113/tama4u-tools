"""VDP+ 한글화 도구: payload 캐시, 번역 대상 추출, 다시 묶기.

  python3 tools/vdp_ko.py cache            # 16개 세트의 payload를 풀어 캐시
  python3 tools/vdp_ko.py extract          # 번역 대상 목록(JSON) 출력
  python3 tools/vdp_ko.py sheet            # 번역 시트 새로 만들기(기존 한국어 칸 유지)
  python3 tools/vdp_ko.py pull             # 시트의 한국어 칸 -> tama4u/vdp-ko/*.json
  python3 tools/vdp_ko.py build circus     # tama4u/vdp-ko/circus.json -> ~/Downloads/TMGC_VDP_ko/

빌드 순서: 영어판 payload에서 시작해
  1. 고정 칸(외출지 이름 8, 캐릭터 이름 14·대사 24, 아이템 이름 10바이트)을 덮어쓰고
  2. 외출지·미니게임 대사는 패킷 끝에 새로 붙인 뒤, 가리키던 코드(xld.w)와
     포인터 표를 새 주소로 바꾸고, 패킷 크기(0x4A)·복사 크기(0x102)·체크섬을 고친다.
  3. payload 안의 내용물은 [u16 길이][패킷][0 채움] 레코드가 4바이트 단위로
     이어진 사슬이라(끝은 FFFF), 커진 패킷의 u16을 고치면 뒤는 그대로 밀린다.
  4. LZ로 다시 압축해 원래 파일 수만큼 나누고 JPEG를 앞에 붙인다.
한글은 한글 패치 펌웨어에서만 보인다(2바이트 글리프).
"""
import glob, os, sys, json, pickle
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))
import struct
from tama4u import container, vdp, charset, outing, sprites, models

ROOT = os.path.expanduser('~/Desktop/Virtual Deco Pierces/vdp_original/')
CACHE = os.environ.get('VDP_CACHE', os.path.expanduser('~/.cache/tama4u-vdp'))


def sets():
    return sorted(glob.glob(ROOT + '*-vdp-plus*'))


def load_set(d):
    """(files, main packet, continuation packets in order)."""
    fs = sorted(glob.glob(d + '/*.jpg'))
    pk = {f: container.parse_file(open(f, 'rb').read())[1][0] for f in fs}
    main = next(f for f, p in pk.items() if vdp.is_vdp(p))
    parts = sorted((f for f, p in pk.items() if vdp.is_part(p)),
                   key=lambda f: vdp.part_index(pk[f]))
    return main, parts, pk


def cached_payload(d):
    name = os.path.basename(d)
    path = os.path.join(CACHE, name + '.bin')
    if os.path.exists(path):
        return open(path, 'rb').read()
    main, parts, pk = load_set(d)
    extra = b''.join(vdp.part_stream(pk[f]) for f in parts)
    data = vdp.payload(pk[main], extra)
    os.makedirs(CACHE, exist_ok=True)
    open(path, 'wb').write(data)
    return data


ROMS = ROOT + 'pierces-roms/'
ROM_OF = {
    'anniversary': 'Anniversary_Change!.bin', 'babychange': 'Baby!_Change!.bin',
    'berrysweets': 'Berry_Sweets_Ver..bin', 'ciao': 'CI\u0e17O_Character_Change!.bin',
    'circus': 'Tama_Star_Circus_Ver..bin',
    'disneychange': 'Disney_Characters_Mickey_&_Friends.bin',
    'fairychange': 'Fairy_Change!.bin', 'lovemelody': 'Love_&_Melody.bin',
    'melodyland': 'Melody_Land_Ver..bin', 'miracrise': 'Miracrise_Ver..bin',
    'nameko': 'Nameko_Nnf_Nnf_Cultivation_Ver.bin', 'royalchange': 'Royal_Change!.bin',
    'sanrio': 'Sanrio_Characters_Mix_Change!.bin', 'yumemiru': 'Yumemiru_Change_Ribbon.bin',
}
PROGRAMS = ('94024702', '94025b01')       # outing, minigame
TAB = None


def short(d):
    return os.path.basename(d).split('-')[0]


def rom_payload(d, n):
    """The pierce ROM's matching bytes (Japanese original), or None."""
    f = ROM_OF.get(short(d))
    if not f:
        return None
    return open(ROMS + f, 'rb').read()[0x1000:0x1000 + n]


def text(codes):
    return charset.decode(list(codes), TAB).rstrip('\u3000')


def cstr(b, o):
    e = o
    while e < len(b) and b[e] != 0:
        e += 1
    return b[o:e]


def sprite_end(it):
    recs = sprites.scan_loose(it)
    return max(r[0] + sprites.loose_span(tuple(r)) for r in recs) if recs else 0x104


def program_strings(it):
    """{key: (offset, bytes)} for every string a program's code points at.

    Strings sit in the data area after the code, usually after the last
    sprite but sometimes between sprites (melodyland, miracrise).  The code
    reaches them either directly (an ``xld.w`` immediate) or through a table
    of LE32 pointers that an immediate points at.  Keys are the instruction
    address (plus the table slot) so the Japanese and English builds pair up."""
    base = outing._detect_base(it)
    recs = sprites.scan_loose(it)
    lo, hi = (recs[0][0] if recs else 0x104), len(it) - 2
    spans = [(r[0], r[0] + sprites.loose_span(tuple(r))) for r in recs]
    ok = lambda fo: lo <= fo < hi and not any(a <= fo < b for a, b in spans)
    # the English releases moved some tables into the spare header bytes
    # (miracrise: 0xA8-0x100) and their strings over sprite-looking bytes,
    # so a table may sit in the header and its entries anywhere in the data
    head = lambda fo: 0x60 <= fo < 0x104
    anywhere = lambda fo: lo <= fo < hi          # data, not code (jump tables)
    out = {}
    refs = list(outing._code_refs(it))
    is_table = lambda fo: (fo + 8 <= hi and (ok(fo) or head(fo)) and
                           all(anywhere(x - base) for x in struct.unpack_from('<II', it, fo)))
    starts = sorted({imm - base for _pc, imm in refs if is_table(imm - base)})
    for pc, imm in refs:
        fo = imm - base
        if not (ok(fo) or head(fo)) or fo + 8 > hi:
            continue
        if is_table(fo):
            # tables sit back to back with no terminator: stop at the next one
            stop = next((t for t in starts if t > fo), hi)
            k = 0
            while fo + 4 * k + 4 <= min(stop, hi):     # ... or at its strings
                v = struct.unpack_from('<I', it, fo + 4 * k)[0]
                if not anywhere(v - base):
                    break
                out[(pc, k)] = (v - base, cstr(it, v - base))
                if fo < v - base < stop:
                    stop = v - base
                k += 1
        elif ok(fo):
            out[(pc, None)] = (fo, cstr(it, fo))
    return out


def packets(data):
    """[(offset, size)] of the contents, by scanning for TAMAGO.

    container.parse_file loops forever on the bare bodies some change
    bundles carry (a header with size 0), so walk it by hand."""
    out, o = [], data.find(container.MAGIC)
    while o >= 0:
        n = struct.unpack_from('>H', data, o + container.OFF_PACKET_SIZE)[0]
        sig = struct.unpack_from('>H', data, o + container.OFF_TYPE_SIG)[0]
        if n >= 0x60 and o + n <= len(data) and models.from_signature(sig):
            out.append((o, n))
            o = data.find(container.MAGIC, o + n)
        else:
            o = data.find(container.MAGIC, o + 6)
    return out


def extract(d):
    data = cached_payload(d)
    jp = rom_payload(d, len(data))
    pk = packets(data)
    rows = []
    row = lambda kind, where, en, ja, cap: rows.append(
        dict(set=short(d), kind=kind, where=where, en=en, ja=ja, cap=cap))
    J = jp or b''
    row('dest', '0x0A', text(data[0x0A:0x12]), text(J[0x0A:0x12]) if jp else '', 8)
    for c in range(data[vdp.CHAR_COUNT_AT]):
        b = vdp.CHAR_BLOCK + c * vdp.CHAR_STRIDE
        o = b + vdp.CH_NAME
        row('char', hex(o), text(data[o:o + 14]), text(J[o:o + 14]) if jp else '', 14)
        for i in range(vdp.CH_LINES_N):
            o = b + vdp.CH_LINES + i * 24
            row('char-line', hex(o), text(data[o:o + 24]), text(J[o:o + 24]) if jp else '', 24)
    for o, n in pk:
        en = data[o:o + n]
        ja = J[o:o + len(en)] if jp else b''
        dest = en[0x4E:0x52].hex()
        row('name:' + dest, hex(o), text(en[0x5E:0x6C]), text(ja[0x5E:0x6C]) if jp else '', 14)
        if dest in PROGRAMS:
            es = program_strings(en)
            js = program_strings(ja) if jp else {}
            for key in sorted(set(es) | set(js), key=lambda k: (k[0], k[1] or 0)):
                e = es.get(key, (0, b''))[1]
                j = js.get(key, (0, b''))[1]
                where = '%s:%x%s' % (hex(o), key[0], '' if key[1] is None else '#%d' % key[1])
                row('text:' + dest, where, text(e), text(j), len(j) or len(e))
    return rows


def cmd_extract():
    global TAB
    TAB = charset.load_table(model="P's")
    allrows = []
    for d in sets():
        rows = extract(d)
        allrows += rows
        print(short(d), len(rows), sum(r['kind'].startswith('text') for r in rows), flush=True)
    out = sys.argv[2] if len(sys.argv) > 2 else os.path.join(CACHE, 'rows.json')
    json.dump(allrows, open(out, 'w'), ensure_ascii=False, indent=0)


def cmd_cache():
    for d in sets():
        data = cached_payload(d)
        print(os.path.basename(d), len(data), flush=True)


KO_DIR = os.path.join(os.path.dirname(__file__), '..', 'tama4u', 'vdp-ko')
OUT_DIR = os.path.expanduser('~/Downloads/TMGC_VDP_ko')
NAME_AT, NAME_BYTES = 0x5E, 10            # zero-terminated; 0x69 is another field
FULL = 0x6C                               # a full-width blank cell


# the download charset's names for symbols the firmware table calls otherwise
# (same codes: 0x66 anger mark, 0x64, 0x62 sun, 0x60)
SYMBOLS = {'!': '！', '?': '？', '╬': '✧', '＠': '◎', '☀': '☼', '☓': '×'}


def enc(t):
    from tama4u import hangul_text
    return hangul_text.encode(''.join(SYMBOLS.get(c, c) for c in t))


def char_line(t):
    """A raising-condition line: 24 bytes shown as two rows of 12 cells.
    Spaces become full cells (1 byte, not the 2-byte half space), and a line
    longer than one row is padded so the second row starts at a word."""
    if '/' not in t and width_cells(t) <= 12:
        return t.replace(' ', '\u3000')
    if '/' in t:
        a, b = t.split('/', 1)
    else:
        words, a = t.split(' '), ''
        while words and width_cells((a + ' ' + words[0]).strip()) <= 12:
            a = (a + ' ' + words.pop(0)).strip()
        b = ' '.join(words)
    a = a.replace(' ', '\u3000')
    return a + '\u3000' * (12 - width_cells(a)) + b.replace(' ', '\u3000')


def width_cells(t):
    return len(t)


def put_fixed(buf, off, cap, t, pad):
    b = enc(t)
    if len(b) > cap:
        raise ValueError(f'{hex(off)} {t!r}: {len(b)}B > {cap}B')
    buf[off:off + cap] = b + bytes([pad]) * (cap - len(b))


def wrap_text(t):
    """12-cell talk lines (11.5: a syllable is a bit wider than a kana).
    A '/' is kept as a forced break and each piece is wrapped on its own;
    a piece that is only a full-width space is a blank line (page push)."""
    import rewrap12
    return '/'.join(p if not p.strip() else rewrap12.wrap(p.strip(), 11.5)
                    for p in t.split('/'))


def patch_program(it, ko):
    """Append the Korean strings and point the code at them."""
    it = bytearray(it)
    old = program_strings(bytes(it))
    base = outing._detect_base(bytes(it))
    end = len(it) - 2                         # drop the checksum, re-add below
    body = it[:end]
    while len(body) % 4:
        body.append(0)
    refs = {pc: imm for pc, imm in outing._code_refs(bytes(it))}
    placed = {}
    for key, (off, _b) in sorted(old.items(), key=lambda kv: kv[1][0]):
        t = ko.get(key)
        if t is None:
            continue
        b = enc(wrap_text(t)) + b'\0'
        if b not in placed:
            placed[b] = len(body)
            body += b
        new = base + placed[b]
        pc, slot = key
        if slot is None:
            outing._reencode_ldw(body, pc, new)
        else:
            struct.pack_into('<I', body, refs[pc] - base + 4 * slot, new)
    while len(body) % 2:
        body.append(0)
    body += b'\0\0'
    if struct.unpack_from('>H', it, 0x102)[0]:
        struct.pack_into('>H', body, 0x102, len(body) - 0x106)
    struct.pack_into('>H', body, container.OFF_PACKET_SIZE, len(body))
    struct.pack_into('>H', body, len(body) - 2, container.sum16(body[:-2]))
    return bytes(body)


def record_len(n):
    """The u16 in front of a packet: its length padded so the record
    ([u16][packet]) stays a multiple of 4."""
    return n + (2 - n) % 4


def build(name):
    d = next(x for x in sets() if short(x) == name)
    ko = json.load(open(os.path.join(KO_DIR, name + '.json')))
    data = bytearray(cached_payload(d))
    if '0x0A' in ko:
        put_fixed(data, 0x0A, 8, ko['0x0A'], FULL)
    for c in range(data[vdp.CHAR_COUNT_AT]):
        b = vdp.CHAR_BLOCK + c * vdp.CHAR_STRIDE
        o = b + vdp.CH_NAME
        if hex(o) in ko:
            put_fixed(data, o, 14, ko[hex(o)], FULL)
        for i in range(vdp.CH_LINES_N):
            o = b + vdp.CH_LINES + i * 24
            if hex(o) in ko:
                put_fixed(data, o, 24, char_line(ko[hex(o)]), FULL)
    # contents, last first so earlier offsets stay valid while we splice
    for o, n in reversed(packets(bytes(data))):
        pkt = bytearray(data[o:o + n])
        if hex(o) in ko:
            if len(enc(ko[hex(o)])) > NAME_BYTES:
                raise ValueError(f'{hex(o)} {ko[hex(o)]!r}: 이름은 {NAME_BYTES}바이트까지')
            put_fixed(pkt, NAME_AT, NAME_BYTES + 1, ko[hex(o)], 0)
        dest = pkt[0x4E:0x52].hex()
        if dest in PROGRAMS:
            mine = {}
            for k, v in ko.items():
                if k.startswith(hex(o) + ':'):
                    w = k.split(':')[1]
                    pc, _, slot = w.partition('#')
                    mine[(int(pc, 16), int(slot) if slot else None)] = v
            if mine:
                pkt = bytearray(patch_program(pkt, mine))
        struct.pack_into('>H', pkt, len(pkt) - 2, container.sum16(pkt[:-2]))
        if len(pkt) != n:
            was = struct.unpack_from('>H', data, o - 2)[0]
            if was != record_len(n):
                raise ValueError(f'{hex(o)}: 사슬 밖 패킷은 길이를 못 바꿉니다')
            rec = record_len(len(pkt))
            data[o - 2:o + was] = struct.pack('>H', rec) + pkt + bytes(rec - len(pkt))
        else:
            data[o:o + n] = pkt
    return d, bytes(data)


def cmd_build():
    name = sys.argv[2]
    d, data = build(name)
    main, parts, pk = load_set(d)
    blobs = vdp.repack_split(pk[main], [pk[f] for f in parts], data)
    out = os.path.join(OUT_DIR, name)
    os.makedirs(out, exist_ok=True)
    for f, blob in zip([main] + parts, blobs):
        jpeg = container.parse_file(open(f, 'rb').read())[0]
        open(os.path.join(out, os.path.basename(f)), 'wb').write(jpeg + blob)
        print(os.path.basename(f), len(blob))
    open(os.path.join(CACHE, name + '-ko.bin'), 'wb').write(data)


SHEET = os.path.expanduser('~/Downloads/TMGC_VDP_번역시트.xlsx')
KIND = {'dest': '외출지 이름', 'char': '캐릭터 이름', 'char-line': '캐릭터 대사',
        'text:94024702': '외출지 대사', 'text:94025b01': '미니게임 대사'}
ITEM = {'81033300': '캐릭터/옷', '81071f00': '거실', '81010101': '로딩 아이콘',
        '81010201': '식사', '81010d02': '간식', '81042a01': '장난감',
        '81042902': '아이콘 세트', '81023400': '액세서리', '94025b01': '미니게임',
        '94024702': '외출지'}
CAP = {'dest': 8, 'char': 14, 'char-line': 24}
HEAD = ['세트', '종류', '키', '일본어 원문', '영어(VDP+)', '한도', '한국어', '바이트', '비고']


def kind_label(k):
    if k.startswith('name:'):
        return '아이템 이름(%s)' % ITEM.get(k[5:], k[5:])
    return KIND.get(k, k)


def nbytes(t, kind=''):
    try:
        if kind == 'char-line':
            return len(enc(char_line(t)))
        return len(enc(wrap_text(t) if kind.startswith('text') else t))
    except Exception:
        return None


def is_text(r):
    """Drop refs that are not dialogue: unchanged by the English release, or
    too short to be a line."""
    if not r['kind'].startswith('text'):
        return True
    return r['ja'] != r['en'] and len(r['ja'] or r['en']) >= 4 if r['ja'] else len(r['en']) >= 4


def cmd_sheet():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    rows = json.load(open(os.path.join(CACHE, 'rows.json')))
    keep = {}
    if os.path.exists(SHEET):
        import shutil, time
        shutil.copy(SHEET, SHEET.replace('.xlsx', '_backup_%s.xlsx' % time.strftime('%m%d%H%M')))
        ws = openpyxl.load_workbook(SHEET)['번역']
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[0] and r[6]:
                keep[(r[0], r[2])] = (r[6], r[8])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '번역'
    ws.append(HEAD)
    font = Font(name='Arial', size=10)
    for c in ws[1]:
        c.font = Font(name='Arial', size=10, bold=True)
        c.fill = PatternFill('solid', fgColor='D9E1F2')
    drafts = {}
    for f in glob.glob(os.path.join(KO_DIR, '*.json')):
        drafts[os.path.basename(f)[:-5]] = json.load(open(f))
    for r in rows:
        if not is_text(r):
            continue
        k = r['kind'].split(':')[0]
        cap = CAP.get(k, NAME_BYTES if k == 'name' else '없음(늘어남)')
        ko, note = keep.get((r['set'], r['where']), (drafts.get(r['set'], {}).get(r['where']), None))
        ws.append([r['set'], kind_label(r['kind']), r['where'], r['ja'], r['en'], cap,
                   ko, nbytes(ko, r['kind']) if ko else None, note])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = font
            c.alignment = Alignment(wrap_text=True, vertical='top')
        row[6].fill = PatternFill('solid', fgColor='FFF2CC')
    for col, w in zip('ABCDEFGHI', (12, 18, 16, 50, 40, 10, 50, 8, 30)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'D2'
    ws.auto_filter.ref = ws.dimensions
    info = wb.create_sheet('설명')
    for line in (
        '노란 칸(한국어)만 고치면 됩니다. pull → build 순서로 반영합니다.',
        '한도: 바이트. 한글·띄어쓰기 2바이트, 기호·영문 1바이트. 이름은 10바이트(한글 5자).',
        '외출지·미니게임 대사는 길이 제한이 없고(패킷이 늘어남) 12칸에서 자동 줄바꿈, /로 직접 줄바꿈.',
        '캐릭터 대사는 24바이트 고정 칸, 한 줄(12칸) 안에 넣기.',
        '바이트 칸은 시트를 만들 때 계산한 값입니다(빌드 때 다시 검사).',
        '키: 이름·대사의 위치(payload 오프셋, 대사는 패킷:명령어주소#표 순번).'):
        info.append([line])
    info.column_dimensions['A'].width = 100
    wb.save(SHEET)
    print(SHEET, ws.max_row - 1)


def cmd_check():
    """Byte limits of every draft against its slot, without building."""
    rows = json.load(open(os.path.join(CACHE, 'rows.json')))
    kinds = {(r['set'], r['where']): r['kind'] for r in rows}
    bad = 0
    for f in sorted(glob.glob(os.path.join(KO_DIR, '*.json'))):
        name = os.path.basename(f)[:-5]
        for k, t in json.load(open(f)).items():
            kind = kinds.get((name, k))
            if kind is None:
                print(name, k, '키 없음'); bad += 1; continue
            kk = kind.split(':')[0]
            cap = CAP.get(kk, NAME_BYTES if kk == 'name' else None)
            try:
                n = nbytes(t, kind)
            except Exception as e:
                print(name, k, t, e); bad += 1; continue
            if n is None:
                print(name, k, repr(t), '인코딩 실패'); bad += 1
            elif cap and n > cap:
                print(name, k, repr(t), f'{n}B > {cap}B'); bad += 1
            elif kk == 'char-line' and '/' not in t and len(t) > 24:
                print(name, k, repr(t), '두 줄 넘음'); bad += 1
    print('문제', bad)


def cmd_pull():
    import openpyxl
    ws = openpyxl.load_workbook(SHEET, data_only=True)['번역']
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] and r[6]:
            out.setdefault(r[0], {})[r[2]] = str(r[6])
    for name, ko in out.items():
        json.dump(ko, open(os.path.join(KO_DIR, name + '.json'), 'w'),
                  ensure_ascii=False, indent=0)
        print(name, len(ko))


if __name__ == '__main__':
    globals()['TAB'] = charset.load_table(model="P's")
    {'cache': cmd_cache, 'extract': cmd_extract, 'build': cmd_build,
     'sheet': cmd_sheet, 'pull': cmd_pull, 'check': cmd_check}[sys.argv[1]]()
