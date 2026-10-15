"""Take the user's redrawn pictures from the sprite sheet's draft column (H)
into tama4u/sprites-ps-ko.json.

    python3 tools/sprite_sheet_import.py ~/Downloads/TMGC_PS_스프라이트시트.xlsx P004 P005 ...

Colours: transparent pixels -> index 0; pure white -> index 0 only where the
ROM original is transparent too (editors leave 255,255,255 specks inside
opaque art, and an index 0 there lets the scene behind show through when
the picture is drawn over another -- P032/P033); exact palette colours keep
the index the current draft has; anything else goes to the nearest opaque
palette colour (counted).
Needs openpyxl + Pillow.
"""
import io, json, os, sys
import openpyxl
from PIL import Image

ROM = open(os.path.expanduser('~/Downloads/TMGC_PS.BIN'), 'rb').read()
JS = os.path.join(os.path.dirname(__file__), '..', 'tama4u', 'sprites-ps-ko.json')


def info(a):
    w, h, nc, _, n, _ = ROM[a:a + 6]; pal = []
    for j in range(nc):
        v = ROM[a + 6 + 2 * j] << 8 | ROM[a + 7 + 2 * j]
        pal.append(((v & 31) << 3, ((v >> 5) & 63) << 2, (v >> 11) << 3))
    return w, h, nc, n, pal


def unpack(raw, w, h, n):
    per = w * h // 2
    return [[(raw[k * per + i // 2] >> 4) if i & 1 else raw[k * per + i // 2] & 15 for i in range(w * h)]
            for k in range(n)]


def decode(img, a, ref, orig=None):
    w, h, nc, n, pal = info(a)
    orig = orig or ref
    sc = img.width // w; px = img.convert('RGBA').load(); out = []; off = 0
    clear = pal[0] == (0, 252, 0)
    for y in range(h):
        for x in range(w):
            r, g, b, al = px[x * sc, y * sc]; c = (r, g, b); o = ref[y * w + x]
            if al < 128 or (clear and c == (255, 255, 255) and orig[y * w + x] == 0):
                out.append(0); continue
            if pal[o] == c:
                out.append(o); continue
            m = [i for i, p in enumerate(pal) if p == c]
            if m:
                out.append(m[0]); continue
            off += 1
            out.append(min(range(1 if clear else 0, nc),
                           key=lambda i: sum((pal[i][j] - c[j]) ** 2 for j in range(3))))
    return out, off


if __name__ == '__main__':
    book, want = os.path.expanduser(sys.argv[1]), set(sys.argv[2:])
    patches = json.load(open(JS))
    wb = openpyxl.load_workbook(book); ws = wb['스프라이트']
    rows = {r: (ws.cell(r, 1).value, ws.cell(r, 2).value) for r in range(2, ws.max_row + 1)}
    for im in ws._images:
        if im.anchor._from.col != 7: continue
        pid, ad = rows[im.anchor._from.row + 1]; base = pid.split('-')[0]
        if base not in want: continue
        k = int(pid.split('-')[1]) - 1 if '-' in pid else 0
        a = int(ad, 16); w, h, nc, n, pal = info(a)
        off0 = a + 6 + 2 * nc
        raw = bytes.fromhex(patches[base]['pixels']) if base in patches else ROM[off0:off0 + n * w * h // 2]
        frames = unpack(raw, w, h, n)
        rom = unpack(ROM[off0:off0 + n * w * h // 2], w, h, n)[k]
        frames[k], off = decode(Image.open(io.BytesIO(im._data())), a, frames[k], rom)
        patches[base] = dict(addr=f'0x{a:06X}', pixels=b''.join(
            bytes(f[i] | f[i + 1] << 4 for i in range(0, len(f), 2)) for f in frames).hex())
        print(pid, 'imported', f'({off} px snapped to the palette)' if off else '')
    json.dump(patches, open(JS, 'w'), indent=0, sort_keys=True)
